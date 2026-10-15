import os
import random
import string
import io
from datetime import datetime, timedelta
from functools import wraps
from flask import Flask, render_template, request, redirect, session, flash, send_file, jsonify
from flask_sqlalchemy import SQLAlchemy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'super_secret_key_change_me')

# База данных: на Render берётся из DATABASE_URL (PostgreSQL), локально - SQLite
db_url = os.environ.get('DATABASE_URL', 'sqlite:///database.db')
# Явно указываем драйвер psycopg2 (он есть в requirements.txt), иначе новые версии
# SQLAlchemy могут искать другой драйвер и сервис не запустится
if db_url.startswith('postgres://'):
    db_url = db_url.replace('postgres://', 'postgresql+psycopg2://', 1)
elif db_url.startswith('postgresql://'):
    db_url = db_url.replace('postgresql://', 'postgresql+psycopg2://', 1)

app.config['SQLALCHEMY_DATABASE_URI'] = db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {'pool_pre_ping': True}

# Чтобы вход работал, когда дневник вставлен в Тильду через iframe,
# cookie должны разрешать работу внутри чужой страницы (нужен HTTPS - на Render он есть)
if os.environ.get('RENDER') or os.environ.get('COOKIE_CROSS_SITE') == '1':
    app.config['SESSION_COOKIE_SAMESITE'] = 'None'
    app.config['SESSION_COOKIE_SECURE'] = True
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=14)

db = SQLAlchemy(app)

# ==================== МОДЕЛИ ====================

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password = db.Column(db.String(80), nullable=False)
    role = db.Column(db.String(20), nullable=False)
    name = db.Column(db.String(100), nullable=True)
    class_id = db.Column(db.Integer, db.ForeignKey('school_class.id'), nullable=True)

class SchoolClass(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(50), unique=True, nullable=False)
    subjects = db.relationship('Subject', backref='school_class', cascade="all, delete-orphan")
    students = db.relationship('Student', backref='school_class', cascade="all, delete-orphan")
    homeworks = db.relationship('Homework', backref='school_class', cascade="all, delete-orphan")
    attendances = db.relationship('Attendance', backref='school_class', cascade="all, delete-orphan")

class Subject(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    class_id = db.Column(db.Integer, db.ForeignKey('school_class.id'), nullable=False)

class Student(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    class_id = db.Column(db.Integer, db.ForeignKey('school_class.id'), nullable=False)
    grades = db.relationship('Grade', backref='student', cascade="all, delete-orphan")

class Grade(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    subject_name = db.Column(db.String(100), nullable=False)
    value = db.Column(db.String(5), nullable=False)
    date = db.Column(db.String(20), nullable=True)
    student_id = db.Column(db.Integer, db.ForeignKey('student.id'), nullable=False)

class Homework(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    date = db.Column(db.String(20), nullable=False)
    subject_name = db.Column(db.String(100), nullable=False)
    task = db.Column(db.Text, nullable=False)
    link = db.Column(db.String(500), nullable=True)
    class_id = db.Column(db.Integer, db.ForeignKey('school_class.id'), nullable=False)
    is_personal = db.Column(db.Boolean, default=False)
    student_name = db.Column(db.String(100), nullable=True)

class HomeworkStatus(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    homework_id = db.Column(db.Integer, db.ForeignKey('homework.id', ondelete='CASCADE'), nullable=False)
    student_name = db.Column(db.String(100), nullable=False)
    status = db.Column(db.String(20), nullable=False)

class Attendance(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    date = db.Column(db.String(20), nullable=False)
    student_name = db.Column(db.String(100), nullable=False)
    status = db.Column(db.String(10), nullable=False)
    class_id = db.Column(db.Integer, db.ForeignKey('school_class.id'), nullable=False)

with app.app_context():
    db.create_all()
    if not User.query.filter_by(username='teacher').first():
        teacher = User(username='teacher', password=os.environ.get('TEACHER_PASSWORD', '1234'), role='teacher', name='Учитель')
        db.session.add(teacher)
        db.session.commit()

# ==================== ВСПОМОГАТЕЛЬНЫЕ ====================

SUBJECT_TEMPLATES = [
    "Математика", "Русский язык", "Литература", "История", "Обществознание",
    "Физика", "Химия", "Биология", "География", "Английский язык",
    "Информатика", "Физкультура", "ОБЖ", "Музыка", "ИЗО"
]

DAY_NAMES = {
    0: "Понедельник", 1: "Вторник", 2: "Среда",
    3: "Четверг", 4: "Пятница", 5: "Суббота", 6: "Воскресенье"
}

def generate_password(length=8):
    chars = string.ascii_letters + string.digits
    return ''.join(random.choice(chars) for _ in range(length))

def get_today_date():
    return datetime.now().strftime("%Y-%m-%d")

def get_monday(d=None):
    if d is None:
        d = datetime.now().date()
    elif isinstance(d, str):
        d = datetime.strptime(d, "%Y-%m-%d").date()
    return d - timedelta(days=d.weekday())

def get_week_info(monday):
    result = []
    for i in range(6):
        day = monday + timedelta(days=i)
        result.append({
            "iso": day.strftime("%Y-%m-%d"),
            "display": day.strftime("%d.%m.%Y"),
            "weekday": DAY_NAMES[day.weekday()]
        })
    return result

def format_classes_dict(week_monday=None):
    if week_monday is None:
        week_monday = get_monday()
    week_info = get_week_info(week_monday)
    week_dates = {d["iso"] for d in week_info}

    result = {}
    for c in SchoolClass.query.all():
        students_list = []
        for st in c.students:
            st_grades = {}
            for g in st.grades:
                st_grades.setdefault(g.subject_name, []).append({
                    "value": g.value,
                    "date": g.date or "",
                    "id": g.id
                })
            students_list.append({"id": st.id, "name": st.name, "grades": st_grades})

        schedule = {}
        personal = {}
        for hw in c.homeworks:
            if hw.date not in week_dates:
                continue
            item = {
                "id": hw.id,
                "date": hw.date,
                "subject": hw.subject_name,
                "task": hw.task,
                "link": hw.link,
                "is_personal": hw.is_personal,
                "student_name": hw.student_name
            }
            if hw.is_personal and hw.student_name:
                personal.setdefault(hw.student_name, []).append(item)
            else:
                schedule.setdefault(hw.date, []).append(item)

        hw_statuses = {}
        for s in HomeworkStatus.query.all():
            hw = Homework.query.get(s.homework_id)
            if hw and hw.class_id == c.id:
                hw_statuses.setdefault(str(hw.id), {})[s.student_name] = s.status

        attendance = {}
        for att in c.attendances:
            attendance.setdefault(att.date, {})[att.student_name] = att.status

        result[c.name] = {
            "id": c.id,
            "students": students_list,
            "subjects": [s.name for s in c.subjects],
            "schedule": schedule,
            "personal_homeworks": personal,
            "homework_status": hw_statuses,
            "attendance": attendance
        }
    return result

def wants_json():
    return request.headers.get('X-Requested-With') == 'XMLHttpRequest'

def teacher_required(f):
    """Все служебные действия доступны только вошедшему учителю."""
    @wraps(f)
    def wrapper(*args, **kwargs):
        user = User.query.filter_by(username=session.get('username')).first() if 'username' in session else None
        if not user or user.role != 'teacher':
            if wants_json():
                return jsonify(success=False, message='Нет доступа'), 403
            return redirect('/login')
        return f(*args, **kwargs)
    return wrapper

# ==================== МАРШРУТЫ ====================

@app.route('/')
def index():
    if 'username' not in session:
        return redirect('/login')
    user = User.query.filter_by(username=session['username']).first()
    if not user:
        session.pop('username', None)
        return redirect('/login')

    week_param = request.args.get('week')
    try:
        monday = get_monday(week_param) if week_param else get_monday()
    except:
        monday = get_monday()

    classes_dict = format_classes_dict(monday)
    week_info = get_week_info(monday)
    week_label = f"{week_info[0]['display']} — {week_info[-1]['display']}"

    if user.role == 'teacher':
        return render_template('teacher_diary.html',
                               classes=classes_dict,
                               selected_class=None,
                               get_today_date=get_today_date,
                               week_monday=monday.strftime("%Y-%m-%d"),
                               week_info=week_info,
                               week_label=week_label,
                               subject_templates=SUBJECT_TEMPLATES)

    s_class = SchoolClass.query.get(user.class_id) if user.class_id else None
    if not s_class or s_class.name not in classes_dict:
        return "Класс не найден", 404

    class_data = classes_dict[s_class.name]
    student_data = next((s for s in class_data["students"] if s['name'].lower() == user.name.lower()), None)
    if not student_data:
        return "Ученик не найден", 404

    return render_template('student_diary.html',
                           student=student_data,
                           class_name=s_class.name,
                           class_data=class_data,
                           schedule=class_data.get("schedule", {}),
                           personal=class_data.get("personal_homeworks", {}).get(student_data["name"], []),
                           week_info=week_info,
                           week_label=week_label,
                           week_monday=monday.strftime("%Y-%m-%d"))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        # убираем случайные пробелы и не различаем большие/маленькие буквы в логине
        # (телефоны часто сами делают первую букву заглавной)
        username = (request.form.get('username') or '').strip()
        password = (request.form.get('password') or '').strip()
        user = User.query.filter(db.func.lower(User.username) == username.lower(), User.password == password).first()
        if user:
            session['username'] = user.username
            session.permanent = True
            return redirect('/')
        return render_template('login.html', error="Неверный логин или пароль")
    return render_template('login.html', error=None)

@app.route('/logout')
def logout():
    session.pop('username', None)
    return redirect('/login')

@app.route('/select_class/<class_name>')
@teacher_required
def select_class(class_name):
    if 'username' not in session:
        return redirect('/login')
    week_param = request.args.get('week')
    try:
        monday = get_monday(week_param) if week_param else get_monday()
    except:
        monday = get_monday()

    classes_dict = format_classes_dict(monday)
    if class_name not in classes_dict:
        return "Класс не найден", 404

    week_info = get_week_info(monday)
    week_label = f"{week_info[0]['display']} — {week_info[-1]['display']}"

    return render_template('teacher_diary.html',
                           classes=classes_dict,
                           selected_class=class_name,
                           get_today_date=get_today_date,
                           week_monday=monday.strftime("%Y-%m-%d"),
                           week_info=week_info,
                           week_label=week_label,
                           subject_templates=SUBJECT_TEMPLATES)

@app.route('/add_class', methods=['POST'])
@teacher_required
def add_class():
    class_name = request.form.get('class_name', '').strip()
    if class_name:
        if SchoolClass.query.filter_by(name=class_name).first():
            msg, cat = f"Класс {class_name} уже существует!", "error"
        else:
            db.session.add(SchoolClass(name=class_name))
            db.session.commit()
            msg, cat = f"Класс {class_name} создан!", "success"
        if wants_json():
            return jsonify(success=True, message=msg, category=cat)
        flash(msg, cat)
    return redirect('/')

@app.route('/delete_class/<class_name>', methods=['POST', 'GET'])
@teacher_required
def delete_class(class_name):
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    if s_class:
        User.query.filter_by(class_id=s_class.id).delete()
        db.session.delete(s_class)
        db.session.commit()
        if wants_json():
            return jsonify(success=True)
    return redirect('/')

@app.route('/add_subject/<class_name>', methods=['POST'])
@teacher_required
def add_subject(class_name):
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    subject_name = request.form.get('subject', '').strip()
    if s_class and subject_name:
        if not Subject.query.filter_by(name=subject_name, class_id=s_class.id).first():
            db.session.add(Subject(name=subject_name, class_id=s_class.id))
            db.session.commit()
            if wants_json():
                return jsonify(success=True, subject=subject_name)
    return redirect('/select_class/' + class_name)

@app.route('/add_template_subjects/<class_name>', methods=['POST'])
@teacher_required
def add_template_subjects(class_name):
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    added = []
    if s_class:
        for name in SUBJECT_TEMPLATES:
            if not Subject.query.filter_by(name=name, class_id=s_class.id).first():
                db.session.add(Subject(name=name, class_id=s_class.id))
                added.append(name)
        db.session.commit()
    if wants_json():
        return jsonify(success=True, added=added)
    return redirect('/select_class/' + class_name)

@app.route('/delete_subject/<class_name>/<subject>', methods=['POST', 'GET'])
@teacher_required
def delete_subject(class_name, subject):
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    if s_class:
        subj = Subject.query.filter_by(name=subject, class_id=s_class.id).first()
        if subj:
            db.session.delete(subj)
            db.session.commit()
            if wants_json():
                return jsonify(success=True)
    return redirect('/select_class/' + class_name)

@app.route('/add_student', methods=['POST'])
@teacher_required
def add_student():
    name = request.form.get('name', '').strip()
    class_name = request.form.get('class_name')
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    if s_class and name:
        translit = {'а':'a','б':'b','в':'v','г':'g','д':'d','е':'e','ё':'e','ж':'zh','з':'z','и':'i','й':'y','к':'k','л':'l','м':'m','н':'n','о':'o','п':'p','р':'r','с':'s','т':'t','у':'u','ф':'f','х':'kh','ц':'ts','ч':'ch','ш':'sh','щ':'shch','ъ':'','ы':'y','ь':'','э':'e','ю':'yu','я':'ya'}
        login = ''.join(translit.get(ch, ch) for ch in name.lower())
        if not login or login.isdigit():
            login = f"user_{random.randint(1000, 9999)}"
        if User.query.filter_by(username=login).first():
            login = f"{login}_{random.randint(100, 999)}"
        password = generate_password()
        student = Student(name=name, class_id=s_class.id)
        db.session.add(student)
        db.session.commit()
        user = User(username=login, password=password, role='student', name=name, class_id=s_class.id)
        db.session.add(user)
        db.session.commit()
        if wants_json():
            return jsonify(success=True, name=name, login=login, password=password)
        flash(f"Ученик {name} добавлен! Логин: {login}, Пароль: {password}", "success")
    return redirect('/select_class/' + class_name)

@app.route('/delete_student/<class_name>/<student_name>', methods=['POST', 'GET'])
@teacher_required
def delete_student(class_name, student_name):
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    if s_class:
        st = Student.query.filter_by(name=student_name, class_id=s_class.id).first()
        if st:
            User.query.filter_by(name=student_name, class_id=s_class.id).delete()
            db.session.delete(st)
            db.session.commit()
            if wants_json():
                return jsonify(success=True)
    return redirect('/select_class/' + class_name)

@app.route('/get_credentials/<class_name>/<student_name>')
@teacher_required
def get_credentials(class_name, student_name):
    if 'username' not in session:
        return jsonify(success=False), 401
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    if not s_class:
        return jsonify(success=False)
    user = User.query.filter_by(name=student_name, class_id=s_class.id, role='student').first()
    if user:
        return jsonify(success=True, login=user.username, password=user.password)
    return jsonify(success=False)

@app.route('/add_grade', methods=['POST'])
@teacher_required
def add_grade():
    class_name = request.form.get('class_name')
    student_name = request.form.get('student_name')
    subject = request.form.get('subject')
    grade_val = request.form.get('grade')
    grade_date = request.form.get('grade_date') or get_today_date()
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    if s_class and grade_val and subject:
        st = Student.query.filter_by(name=student_name, class_id=s_class.id).first()
        if st:
            g = Grade(subject_name=subject, value=str(grade_val), date=grade_date, student_id=st.id)
            db.session.add(g)
            db.session.commit()
            if wants_json():
                return jsonify(success=True, grade={"value": g.value, "date": g.date, "id": g.id, "subject": subject})
    return redirect('/select_class/' + class_name)

@app.route('/delete_grade/<int:grade_id>', methods=['POST', 'GET'])
@teacher_required
def delete_grade_by_id(grade_id):
    g = Grade.query.get(grade_id)
    if g:
        db.session.delete(g)
        db.session.commit()
        if wants_json():
            return jsonify(success=True)
    return redirect(request.referrer or '/')

@app.route('/add_homework', methods=['POST'])
@teacher_required
def add_homework():
    class_name = request.form.get('class_name')
    date = request.form.get('date')
    subject = request.form.get('subject')
    task = request.form.get('task')
    link = request.form.get('link', '').strip() or None
    is_personal = request.form.get('is_personal') == '1'
    student_name = request.form.get('student_name') if is_personal else None
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    if s_class and date and subject and task:
        hw = Homework(date=date, subject_name=subject, task=task, link=link,
                      class_id=s_class.id, is_personal=is_personal, student_name=student_name)
        db.session.add(hw)
        db.session.commit()
        if wants_json():
            return jsonify(success=True, homework={
                "id": hw.id, "date": hw.date, "subject": hw.subject_name,
                "task": hw.task, "link": hw.link, "is_personal": hw.is_personal,
                "student_name": hw.student_name
            })
        flash(f"ДЗ по {subject} добавлено!", "success")
    return redirect('/select_class/' + class_name + '?week=' + (date or ''))

@app.route('/delete_homework/<int:hw_id>', methods=['POST', 'GET'])
@teacher_required
def delete_homework(hw_id):
    hw = Homework.query.get(hw_id)
    if hw:
        HomeworkStatus.query.filter_by(homework_id=hw.id).delete()
        db.session.delete(hw)
        db.session.commit()
        if wants_json():
            return jsonify(success=True)
    return redirect(request.referrer or '/')

@app.route('/clear_homeworks/<class_name>', methods=['POST'])
@teacher_required
def clear_homeworks(class_name):
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    week = request.form.get('week') or request.args.get('week')
    if s_class:
        if week:
            monday = get_monday(week)
            week_dates = {d["iso"] for d in get_week_info(monday)}
            for hw in list(s_class.homeworks):
                if hw.date in week_dates:
                    HomeworkStatus.query.filter_by(homework_id=hw.id).delete()
                    db.session.delete(hw)
        else:
            for hw in list(s_class.homeworks):
                HomeworkStatus.query.filter_by(homework_id=hw.id).delete()
                db.session.delete(hw)
        db.session.commit()
        if wants_json():
            return jsonify(success=True)
    return redirect('/select_class/' + class_name + ('?week=' + week if week else ''))

@app.route('/mark_homework/<int:hw_id>/<student_name>', methods=['POST'])
@teacher_required
def mark_homework(hw_id, student_name):
    status_val = request.form.get('status')
    hw = Homework.query.get(hw_id)
    if hw and status_val:
        st_status = HomeworkStatus.query.filter_by(homework_id=hw.id, student_name=student_name).first()
        if st_status:
            st_status.status = status_val
        else:
            db.session.add(HomeworkStatus(homework_id=hw.id, student_name=student_name, status=status_val))
        db.session.commit()
        if wants_json():
            return jsonify(success=True, status=status_val)
    return redirect(request.referrer or '/')

@app.route('/mark_attendance/<class_name>', methods=['POST'])
@teacher_required
def mark_attendance(class_name):
    date = request.form.get('date')
    status = request.form.get('status')
    student_name = request.form.get('student_name')
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    if s_class and date and status and student_name:
        att = Attendance.query.filter_by(class_id=s_class.id, date=date, student_name=student_name).first()
        if att:
            att.status = status
        else:
            db.session.add(Attendance(class_id=s_class.id, date=date, student_name=student_name, status=status))
        db.session.commit()
        if wants_json():
            return jsonify(success=True)
    return redirect('/select_class/' + class_name)

@app.route('/delete_attendance/<class_name>/<date>', methods=['POST', 'GET'])
@teacher_required
def delete_attendance(class_name, date):
    s_class = SchoolClass.query.filter_by(name=class_name).first()
    if s_class:
        Attendance.query.filter_by(class_id=s_class.id, date=date).delete()
        db.session.commit()
        if wants_json():
            return jsonify(success=True)
    return redirect('/select_class/' + class_name)

@app.route('/export_student_stats/<class_name>/<student_name>', methods=['GET', 'POST'])
@teacher_required
def export_student_stats(class_name, student_name):
    classes_dict = format_classes_dict(get_monday())
    if class_name not in classes_dict:
        return "Класс не найден", 404
    student_data = next((s for s in classes_dict[class_name]["students"] if s['name'] == student_name), None)
    if not student_data:
        return "Ученик не найден", 404

    start_date = request.form.get('start_date') or request.args.get('start_date') or '2000-01-01'
    end_date = request.form.get('end_date') or request.args.get('end_date') or '2099-12-31'

    doc = Document()
    title = doc.add_heading('ОТЧЕТ ОБ УСПЕВАЕМОСТИ И ПОСЕЩАЕМОСТИ', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Ученик: {student_name}")
    doc.add_paragraph(f"Класс: {class_name}")
    doc.add_paragraph(f"Период: {start_date} — {end_date}")
    doc.add_paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

    doc.add_heading('Успеваемость по предметам', level=2)
    grades = student_data.get('grades', {})
    has_grades = False
    if grades:
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = 'Предмет', 'Оценки', 'Даты', 'Средний'
        for subject, subj_grades in grades.items():
            filtered = [g for g in subj_grades if start_date <= (g.get('date') or '2000-01-01') <= end_date]
            if filtered:
                has_grades = True
                row = table.add_row().cells
                row[0].text = subject
                vals = [g['value'] for g in filtered]
                row[1].text = ", ".join(vals)
                row[2].text = ", ".join(g.get('date') or '-' for g in filtered)
                numeric = [float(v) for v in vals if v.replace('.', '', 1).isdigit()]
                row[3].text = f"{sum(numeric)/len(numeric):.2f}" if numeric else "-"
    if not has_grades:
        doc.add_paragraph("Оценок за выбранный период нет.")

    doc.add_heading('Посещаемость', level=2)
    attendance = classes_dict[class_name].get('attendance', {})
    records = [(date, record[student_name]) for date, record in sorted(attendance.items())
               if start_date <= date <= end_date and student_name in record]
    if records:
        att_table = doc.add_table(rows=1, cols=2)
        att_table.style = 'Table Grid'
        att_table.rows[0].cells[0].text = 'Дата'
        att_table.rows[0].cells[1].text = 'Статус'
        for d, st in records:
            r = att_table.add_row().cells
            r[0].text = d
            r[1].text = st
    else:
        doc.add_paragraph("Записей о посещаемости за период нет.")

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return send_file(file_stream, as_attachment=True, download_name=f"stats_{student_name}.docx",
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

if __name__ == '__main__':
    import socket
    import os
    hostname = socket.gethostname()
    local_ip = socket.gethostbyname(hostname)
    port = int(os.environ.get('PORT', 5000))
    print("\n" + "="*50)
    print("🚀 ЭЛЕКТРОННЫЙ ДНЕВНИК ЗАПУЩЕН!")
    print("="*50)
    print(f"📍 Доступ: http://0.0.0.0:{port}")
    if local_ip != '127.0.0.1':
        print(f"🌐 Доступ в сети: http://{local_ip}:{port}")
    print("="*50)
    app.run(host='0.0.0.0', port=port, debug=False)
    
